import streamlit as st
import pandas as pd
import calendar
from datetime import datetime, timedelta
from io import BytesIO
import os
import openpyxl
from openpyxl.cell.cell import MergedCell

st.set_page_config(page_title="Plataforma Operacional - Zona 4 Dunkin", layout="wide")

# ==========================================
# 🎨 ESTILOS CSS PERSONALIZADOS (Diseño ERP Unificado)
# ==========================================
st.markdown("""
    <style>
    /* Fondo general blanco y limpio para toda la aplicación */
    .stApp {
        background-color: #ffffff;
        color: #2b2b2b;
    }

    /* Forzar textos generales a gris oscuro/negro para legibilidad en fondo blanco */
    h1, h2, h3, h4, h5, h6, p, span, label {
        color: #2b2b2b !important;
    }

    /* Tarjetas ERP unificadas con altura uniforme para mantener simetría */
    .erp-card {
        background: #f0fbfc;
        backdrop-filter: blur(10px);
        border: 1.5px solid #00b4d8;
        border-radius: 16px;
        padding: 22px;
        color: #2b2b2b;
        text-align: center;
        box-shadow: 0 4px 15px rgba(0, 180, 216, 0.15);
        margin-bottom: 15px;
        min-height: 165px;
        display: flex;
        flex-direction: column;
        justify-content: center;
        align-items: center;
        transition: 0.3s;
    }
    .erp-card:hover {
        border-color: #0077b6;
        box-shadow: 0 6px 20px rgba(0, 119, 182, 0.25);
    }
    .erp-title {
        font-size: 15px;
        font-weight: 700;
        color: #0077b6 !important;
        margin-bottom: 8px;
    }
    .erp-value {
        font-size: 24px;
        font-weight: bold;
        color: #03045e !important;
    }

    /* Contenedor de Login Estilizado */
    .login-container {
        background: #f0fbfc;
        backdrop-filter: blur(10px);
        border: 2px solid #00b4d8;
        border-radius: 20px;
        padding: 40px;
        box-shadow: 0 8px 32px 0 rgba(0, 180, 216, 0.2);
        max-width: 600px;
        margin: 50px auto;
        color: #2b2b2b;
    }
    
    .login-title {
        font-size: 28px;
        font-weight: 800;
        color: #03045e !important;
        text-align: center;
        margin-bottom: 10px;
    }
    
    .login-subtitle {
        font-size: 15px;
        color: #0077b6 !important;
        text-align: center;
        margin-bottom: 30px;
    }

    /* Etiquetas de inputs y selectores de Streamlit */
    .stTextInput label, .stSelectbox label, .stDateInput label {
        color: #03045e !important;
        font-weight: 600 !important;
    }

    /* Estilo robusto para forzar bordes cian en todos los selectores / dropdowns de Streamlit */
    div[data-baseweb="select"] > div {
        background-color: #ffffff !important;
        border: 2px solid #00b4d8 !important;
        border-radius: 8px !important;
        color: #2b2b2b !important;
    }
    div[data-baseweb="select"] > div:hover {
        border-color: #0077b6 !important;
    }

    /* Estilo para los inputs de texto */
    .stTextInput input {
        background-color: #ffffff !important;
        color: #2b2b2b !important;
        border: 2px solid #00b4d8 !important;
        border-radius: 8px !important;
    }
    .stTextInput input:focus {
        border-color: #0077b6 !important;
        box-shadow: 0 0 8px rgba(0, 180, 216, 0.4) !important;
    }

    /* Botones con estilo cian */
    div.stButton > button {
        background-color: #00b4d8 !important;
        color: white !important;
        border: none !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
    }
    div.stButton > button:hover {
        background-color: #0077b6 !important;
        color: white !important;
    }

    /* 📊 Compactar tablas y ocultar por completo la primera columna de números (índice) */
    [data-testid="stDataFrame"] th, [data-testid="stDataFrame"] td {
        font-size: 13px !important;
        padding: 4px 6px !important;
    }
    [data-testid="stDataFrame"] table tr th:first-child,
    [data-testid="stDataFrame"] table tr td:first-child {
        display: none !important;
    }
    </style>
""", unsafe_allow_html=True)

# ==========================================
# 🔐 1. CREDENCIALES Y DATOS INICIALES
# ==========================================
CEDULA_MASTER = "1032463775"      
CLAVE_MASTER = "Carolina2026"  

if "db_admins" not in st.session_state:
    st.session_state.db_admins = {
        "CEDULA_ADMIN_1": {"nombre": "Camilo", "tienda": "PV 102 - CR 12 89 33 LC 1"},
        "CEDULA_ADMIN_2": {"nombre": "Karen", "tienda": "PV 106 - CL 140 12 51 BQ A"},
        "CEDULA_ADMIN_3": {"nombre": "Esteban", "tienda": "PV 136 - EXITO COUNTRY CL 134"},
        "CEDULA_ADMIN_4": {"nombre": "Jennifer", "tienda": "PV 141 - EDS TERPEL AV 13 125"},
        "CEDULA_ADMIN_5": {"nombre": "Carolina", "tienda": "PV 180 - CC UNICENTRO LC2-158"},
        "CEDULA_ADMIN_6": {"nombre": "Danna", "tienda": "PV 225 - EDIF EPSON CL 100"},
        "CEDULA_ADMIN_7": {"nombre": "Everley", "tienda": "PV 226 - EDIF TELEPORT CL 113"},
        "CEDULA_ADMIN_8": {"nombre": "Julian", "tienda": "PV 244 - EDIF WORLD TRADE CENT"},
        "CEDULA_ADMIN_9": {"nombre": "KarenN", "tienda": "PV 249 - EDIF FLORMORADO"},
    }

if "lista_almacenes_base" not in st.session_state:
    st.session_state.lista_almacenes_base = [
        "PV 102 - CR 12 89 33 LC 1", "PV 106 - CL 140 12 51 BQ A",
        "PV 136 - EXITO COUNTRY CL 134", "PV 141 - EDS TERPEL AV 13 125",
        "PV 180 - CC UNICENTRO LC2-158", "PV 225 - EDIF EPSON CL 100",
        "PV 226 - EDIF TELEPORT CL 113", "PV 244 - EDIF WORLD TRADE CENT",
        "PV 249 - EDIF FLORMORADO", "PV 262 - AUTOP NORTE",
        "PV 266 - CR 15 95 55", "PV 276 - CR 15 92 29 LC2",
        "PV 257 - ALCALA", "PV 283 - Au Ed Sa 123-24 L10",
        "PV 284 - Edif Tokio LC1", "PV 285 - TORR EMPRESARIAL FD-100",
        "PV 290 - CL 90 16 10 LC 5 ED", "PV 295 - CALLE 147",
        "PV 296 - NIZA CALLE 127"
    ]

# ==========================================
# 💾 PERSISTENCIA CON ARCHIVOS CSV Y EXCEL EOR
# ==========================================
DIR_DATOS = "datos_persistencia"
if not os.path.exists(DIR_DATOS):
    os.makedirs(DIR_DATOS)

EXCEL_EOR_MAESTRO = "EOR SEPTIEMBRE 2026.xlsx"

def obtener_o_crear_presupuesto(anio, mes_num):
    clave_mes = f"{anio}-{mes_num:02d}"
    ruta_csv = os.path.join(DIR_DATOS, f"presupuesto_{clave_mes}.csv")
    col_venta_pasada = f"Venta Mes {anio - 1} ($)"
    
    if os.path.exists(ruta_csv):
        df_existente = pd.read_csv(ruta_csv)
        if col_venta_pasada not in df_existente.columns:
            col_vieja = [c for c in df_existente.columns if "Venta Mes" in c and str(anio - 1) in c]
            if col_vieja:
                df_existente.rename(columns={col_vieja[0]: col_venta_pasada}, inplace=True)
            else:
                df_existente[col_venta_pasada] = 0.0
            df_existente.to_csv(ruta_csv, index=False)
        return df_existente
    else:
        df_nuevo = pd.DataFrame({
            "Almacén": st.session_state.lista_almacenes_base,
            "Presupuesto Mes ($)": [15000000.0] * len(st.session_state.lista_almacenes_base),
            col_venta_pasada: [0.0] * len(st.session_state.lista_almacenes_base)
        })
        df_nuevo.to_csv(ruta_csv, index=False)
        return df_nuevo

def guardar_presupuesto(anio, mes_num, df):
    clave_mes = f"{anio}-{mes_num:02d}"
    ruta_csv = os.path.join(DIR_DATOS, f"presupuesto_{clave_mes}.csv")
    df.to_csv(ruta_csv, index=False)

def obtener_o_crear_mes(anio, mes_num):
    clave_mes = f"{anio}-{mes_num:02d}"
    ruta_csv = os.path.join(DIR_DATOS, f"registros_{clave_mes}.csv")
    
    if os.path.exists(ruta_csv):
        return pd.read_csv(ruta_csv)
    else:
        _, num_dias = calendar.monthrange(anio, mes_num)
        dias_espanol = {
            'Monday': 'Lunes', 'Tuesday': 'Martes', 'Wednesday': 'Miércoles', 
            'Thursday': 'Jueves', 'Friday': 'Viernes', 'Saturday': 'Sábado', 'Sunday': 'Domingo'
        }
        
        registros = []
        for tienda in st.session_state.lista_almacenes_base:
            for d in range(1, num_dias + 1):
                fecha_obj = datetime(anio, mes_num, d)
                dia_ingles = fecha_obj.strftime('%A')
                dia_esp = dias_espanol.get(dia_ingles, dia_ingles)
                fecha_str = f"{fecha_obj.strftime('%Y-%m-%d')} - {dia_esp}"
                
                registros.append({
                    "Almacén": tienda,
                    "Fecha y Día": fecha_str,
                    "Venta Diaria ($)": 0.0,
                    "Transacciones / Clientes": 0,
                    "Unidades Vendidas": 0,
                    "Desperdicio (Unid)": 0
                })
        df_nuevo = pd.DataFrame(registros)
        df_nuevo.to_csv(ruta_csv, index=False)
        return df_nuevo

def guardar_mes(anio, mes_num, df):
    clave_mes = f"{anio}-{mes_num:02d}"
    ruta_csv = os.path.join(DIR_DATOS, f"registros_{clave_mes}.csv")
    df.to_csv(ruta_csv, index=False)

def escribir_celda_segura(sheet, coordenada, valor):
    """Evita errores de celdas combinadas (MergedCell) asignando al rango principal."""
    celda = sheet[coordenada]
    if isinstance(celda, MergedCell):
        for rango in sheet.merged_cells.ranges:
            if coordenada in rango:
                min_row, min_col, _, _ = rango.bounds
                sheet.cell(row=min_row, column=min_col, value=valor)
                return
    else:
        sheet[coordenada] = valor

if "modulo_actual" not in st.session_state:
    st.session_state.modulo_actual = "Inicio"

# ==========================================
# 🔐 2. SISTEMA DE LOGIN 
# ==========================================
if "autenticado" not in st.session_state:
    st.session_state.autenticado = False
    st.session_state.usuario_rol = None
    st.session_state.cedula_actual = None

def verificar_credenciales(usuario, password):
    if usuario == CEDULA_MASTER and password == CLAVE_MASTER:
        return "Master", None
    elif usuario in st.session_state.db_admins:
        nombre_admin = st.session_state.db_admins[usuario]["nombre"]
        if password == f"{nombre_admin}*":
            return "Administrador", usuario
    return None, None

if not st.session_state.autenticado:
    st.markdown("<br><br>", unsafe_allow_html=True)
    st.markdown("""
        <div class="login-container">
            <div class="login-title">🍩 Plataforma Operacional Zona 4</div>
            <div class="login-subtitle">🔒 Acceso Seguro — Ingresa tus credenciales para continuar</div>
    """, unsafe_allow_html=True)
    
    with st.form("login_form"):
        usuario_input = st.text_input("Usuario (Número de Cédula):")
        password_input = st.text_input("Contraseña:", type="password")
        st.markdown("<br>", unsafe_allow_html=True)
        submitted = st.form_submit_button("Ingresar a la Plataforma", use_container_width=True)
        
        if submitted:
            rol, cedula = verificar_credenciales(usuario_input, password_input)
            if rol:
                st.session_state.autenticado = True
                st.session_state.usuario_rol = rol
                st.session_state.cedula_actual = cedula
                st.rerun()
            else:
                st.error("Cédula o contraseña incorrectos.")
                
    st.markdown("</div>", unsafe_allow_html=True)
    st.stop()

# ==========================================
# 🖥️ 3. BARRA LATERAL (NAVEGACIÓN)
# ==========================================
if st.session_state.usuario_rol == "Master":
    nombre_mostrar = "Carolina Zamora (Consultora)"
else:
    nombre_mostrar = st.session_state.db_admins[st.session_state.cedula_actual]["nombre"]

st.sidebar.markdown(f"👋 **¡Bienvenido(a), {nombre_mostrar}!**")

if st.session_state.modulo_actual != "Inicio":
    if st.sidebar.button("🏠 Menú Principal"):
        st.session_state.modulo_actual = "Inicio"
        st.rerun()

st.sidebar.markdown("---")
if st.sidebar.button("Cerrar Sesión"):
    st.session_state.autenticado = False
    st.session_state.usuario_rol = None
    st.session_state.cedula_actual = None
    st.session_state.modulo_actual = "Inicio"
    st.rerun()

# ==========================================
# 🗂️ 4. ENRUTADOR DE MÓDULOS (ESTILO ERP UNIFICADO)
# ==========================================

if st.session_state.modulo_actual == "Inicio":
    st.title("🍩 Plataforma Operacional - Zona 4 Dunkin")
    st.write("Selecciona el módulo de gestión que deseas operar:")
    st.markdown("<br>", unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns(3)
    
    with col1:
        st.markdown("""
            <div class="erp-card">
                <div class="erp-title">💰 FINANZAS</div>
                <div style="font-size: 13px; color: #555; margin-bottom: 15px;">Presupuesto, Ventas & Desperdicio</div>
            </div>
        """, unsafe_allow_html=True)
        if st.button("🚀 Ingresar a Finanzas", use_container_width=True):
            st.session_state.modulo_actual = "Finanzas"
            st.rerun()
            
        st.markdown("<br>", unsafe_allow_html=True)
        
        st.markdown("""
            <div class="erp-card">
                <div class="erp-title">📦 INVENTARIOS</div>
                <div style="font-size: 13px; color: #555; margin-bottom: 15px;">Control de stock y pedidos</div>
            </div>
        """, unsafe_allow_html=True)
        if st.button("🚀 Ingresar a Inventarios", use_container_width=True):
            st.session_state.modulo_actual = "Inventarios"
            st.rerun()

    with col2:
        st.markdown("""
            <div class="erp-card">
                <div class="erp-title">📅 PLANILLA SEMANAL</div>
                <div style="font-size: 13px; color: #555; margin-bottom: 15px;">Turnos, HeadCount & Rotación Admins</div>
            </div>
        """, unsafe_allow_html=True)
        if st.button("🚀 Ingresar a Planilla", use_container_width=True):
            st.session_state.modulo_actual = "Planilla"
            st.rerun()
            
        st.markdown("<br>", unsafe_allow_html=True)
        
        st.markdown("""
            <div class="erp-card">
                <div class="erp-title">📋 AUDITORÍA EOR</div>
                <div style="font-size: 13px; color: #555; margin-bottom: 15px;">Estándares operativos de calidad</div>
            </div>
        """, unsafe_allow_html=True)
        if st.button("🚀 Ingresar a EOR", use_container_width=True):
            st.session_state.modulo_actual = "EOR"
            st.rerun()

    with col3:
        st.markdown("""
            <div class="erp-card">
                <div class="erp-title">🔮 PREVISIONES</div>
                <div style="font-size: 13px; color: #555; margin-bottom: 15px;">Proyecciones y metas de zona</div>
            </div>
        """, unsafe_allow_html=True)
        if st.button("🚀 Ingresar a Previsiones", use_container_width=True):
            st.session_state.modulo_actual = "Previsiones"
            st.rerun()

# ---> MÓDULO DE FINANZAS <---
elif st.session_state.modulo_actual == "Finanzas":
    st.title("📊 Módulo de Finanzas - Zona 4")
    col_anio, col_mes = st.columns(2)
    with col_anio:
        anio_sel = st.selectbox("Año Actual:", [2025, 2026, 2027, 2028], index=1)
    with col_mes:
        meses_nombres = {
            1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio",
            7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre"
        }
        mes_sel_nombre = st.selectbox("Mes a Consultar / Operar:", list(meses_nombres.values()), index=8)
    
    mes_num = [k for k, v in meses_nombres.items() if v == mes_sel_nombre][0]
    anio_pasado = anio_sel - 1
    col_vp_nombre = f"Venta Mes {anio_pasado} ($)"

    df_mes_activo = obtener_o_crear_mes(anio_sel, mes_num)
    df_presupuesto_activo = obtener_o_crear_presupuesto(anio_sel, mes_num)

    hoy = datetime.now()
    _, total_dias_mes = calendar.monthrange(anio_sel, mes_num)
    
    if anio_sel == hoy.year and mes_num == hoy.month:
        dia_corte_default = max(1, hoy.day - 1)
    elif (anio_sel * 12 + mes_num) < (hoy.year * 12 + hoy.month):
        dia_corte_default = total_dias_mes
    else:
        dia_corte_default = 1

    if st.session_state.usuario_rol == "Master":
        tab_resumen, tab_presupuestos, tab_individual = st.tabs([
            f"📋 Análisis Integral de Almacenes ({mes_sel_nombre})", f"⚙️ Configurar Presupuestos & Venta {anio_pasado} ({mes_sel_nombre})", f"🔍 Registro Diario por Tienda ({mes_sel_nombre})"
        ])
        
        with tab_resumen:
            st.subheader("📋 Análisis Integral de Almacenes")
            primer_dia_mes = datetime(anio_sel, mes_num, 1).date()
            ultimo_dia_mes = datetime(anio_sel, mes_num, total_dias_mes).date()
            fecha_fin_default = datetime(anio_sel, mes_num, dia_corte_default).date() if dia_corte_default <= total_dias_mes else ultimo_dia_mes
            
            st.markdown("### 🗓️ Rango de Fechas")
            col_f1, col_f2 = st.columns(2)
            with col_f1:
                rango_fechas = st.date_input(
                    "Selecciona el rango (Inicio y Fin):",
                    value=(primer_dia_mes, fecha_fin_default),
                    min_value=datetime(2025, 1, 1).date(),
                    max_value=datetime(2028, 12, 31).date(),
                    key=f"rango_fechas_{anio_sel}_{mes_num}"
                )
            
            if isinstance(rango_fechas, tuple) and len(rango_fechas) == 2:
                f_inicio, f_fin = rango_fechas
            elif isinstance(rango_fechas, tuple) and len(rango_fechas) == 1:
                f_inicio = f_fin = rango_fechas[0]
            else:
                f_inicio = primer_dia_mes
                f_fin = fecha_fin_default

            if f_inicio > f_fin:
                f_inicio, f_fin = f_fin, f_inicio
                
            delta_dias = (f_fin - f_inicio).days + 1
            pct_meta_tiempo = (delta_dias / total_dias_mes) * 100

            def filtrar_por_rango(df, inicio, fin):
                df_copia = df.copy()
                df_copia["Fecha_Obj"] = pd.to_datetime(df_copia["Fecha y Día"].apply(lambda x: x.split(" ")[0])).dt.date
                return df_copia[(df_copia["Fecha_Obj"] >= inicio) & (df_copia["Fecha_Obj"] <= fin)]

            df_mes_rango = filtrar_por_rango(df_mes_activo, f_inicio, f_fin)

            def obtener_consolidado_rango(df_rango, df_presup):
                suma_mes = df_rango.groupby("Almacén")[["Venta Diaria ($)", "Transacciones / Clientes", "Unidades Vendidas", "Desperdicio (Unid)"]].sum().reset_index()
                consolidado = pd.merge(df_presup, suma_mes, on="Almacén", how="left").fillna(0)
                consolidado["Ticket Promedio ($)"] = consolidado.apply(
                    lambda row: row["Venta Diaria ($)"] / row["Transacciones / Clientes"] if row["Transacciones / Clientes"] > 0 else 0.0,
                    axis=1
                )
                consolidado["Venta Proporcional Año Pasado ($)"] = consolidado.apply(
                    lambda row: (row[col_vp_nombre] / total_dias_mes) * delta_dias,
                    axis=1
                )
                consolidado = consolidado[["Almacén", "Presupuesto Mes ($)", col_vp_nombre, "Venta Proporcional Año Pasado ($)", "Venta Diaria ($)", "Ticket Promedio ($)", "Unidades Vendidas", "Desperdicio (Unid)"]]
                consolidado.rename(columns={"Venta Diaria ($)": "Venta Acumulada Rango ($)", "Desperdicio (Unid)": "Desperdicio Acumulado (Unid)"}, inplace=True)
                consolidado["Presupuesto Mes ($)"] = consolidado["Presupuesto Mes ($)"].round(0)
                consolidado[col_vp_nombre] = consolidado[col_vp_nombre].round(0)
                consolidado["Venta Proporcional Año Pasado ($)"] = consolidado["Venta Proporcional Año Pasado ($)"].round(0)
                consolidado["Venta Acumulada Rango ($)"] = consolidado["Venta Acumulada Rango ($)"].round(0)
                consolidado["Ticket Promedio ($)"] = consolidado["Ticket Promedio ($)"].round(0)
                consolidado["Desperdicio Acumulado (Unid)"] = consolidado["Desperdicio Acumulado (Unid)"].round(0).astype(int)
                consolidado["Unidades Vendidas"] = consolidado["Unidades Vendidas"].round(0).astype(int)
                return consolidado

            df_consolidado_rango = obtener_consolidado_rango(df_mes_rango, df_presupuesto_activo)

            def calcular_metricas_globales_rango(df_con, df_presup):
                total_venta = df_con["Venta Acumulada Rango ($)"].sum()
                total_presupuesto = df_presup["Presupuesto Mes ($)"].sum()
                cumplimiento = (total_venta / total_presupuesto * 100) if total_presupuesto > 0 else 0
                meta_acumulada_rango = (total_presupuesto / total_dias_mes) * delta_dias
                cumplimiento_proporcional_rango = (total_venta / meta_acumulada_rango * 100) if meta_acumulada_rango > 0 else 0
                total_transacciones = df_mes_rango["Transacciones / Clientes"].sum()
                ticket_prom_zona = (total_venta / total_transacciones) if total_transacciones > 0 else 0.0
                total_unid_zona = df_con["Unidades Vendidas"].sum()
                total_desp_zona = df_con["Desperdicio Acumulado (Unid)"].sum()
                pct_desp_zona = (total_desp_zona / total_unid_zona * 100) if total_unid_zona > 0 else 0.0
                total_venta_prop_pasada = df_con["Venta Proporcional Año Pasado ($)"].sum()
                crecimiento_zona = ((total_venta - total_venta_prop_pasada) / total_venta_prop_pasada * 100) if total_venta_prop_pasada > 0 else 0.0
                return total_venta, cumplimiento, cumplimiento_proporcional_rango, pct_desp_zona, ticket_prom_zona, meta_acumulada_rango, crecimiento_zona

            t_venta, t_cumplimiento, t_cump_prop, t_pct_desp_zona, t_ticket_zona, meta_fecha_zona, crecimiento_zona = calcular_metricas_globales_rango(df_consolidado_rango, df_presupuesto_activo)

            st.markdown(f"### 🌐 Indicadores Globales Zona 4 — {pct_meta_tiempo:.1f}% del mes")
            
            col_k1, col_k2, col_k3, col_k4, col_k5 = st.columns(5)
            with col_k1:
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">💰 VENTA ACUMULADA RANGO</div>
                        <div class="erp-value">${t_venta:,.0f}</div>
                        <div style="font-size: 11px; color: #0077b6; margin-top: 4px;">Meta Proporcional: ${meta_fecha_zona:,.0f}</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_k2:
                color_cump_card = "#28a745" if t_cump_prop >= 100 else ("#ffc107" if t_cump_prop >= 85 else "#dc3545")
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">📈 CUMPLIMIENTO RANGO</div>
                        <div class="erp-value" style="color: {color_cump_card};">{t_cump_prop:.2f}%</div>
                        <div style="font-size: 11px; color: #555; margin-top: 4px;">Cump. Mes Total: {t_cumplimiento:.2f}%</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_k3:
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">🎫 TICKET PROMEDIO ZONA</div>
                        <div class="erp-value">${t_ticket_zona:,.0f}</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_k4:
                color_tarjeta_desp = "#28a745" if t_pct_desp_zona <= 10.0 else "#dc3545"
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">⚠️ % DESPERDICIO ZONA</div>
                        <div class="erp-value" style="color: {color_tarjeta_desp};">{t_pct_desp_zona:.2f}%</div>
                        <div style="font-size: 11px; color: #555; margin-top: 4px;">Límite mes: 10.0%</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_k5:
                color_crec_card = "#28a745" if crecimiento_zona >= 0 else "#dc3545"
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">📊 CRECIMIENTO RANGO</div>
                        <div class="erp-value" style="color: {color_crec_card};">{crecimiento_zona:.2f}%</div>
                        <div style="font-size: 11px; color: #555; margin-top: 4px;">vs {anio_pasado} en rango</div>
                    </div>
                """, unsafe_allow_html=True)
                
            st.markdown("<br>", unsafe_allow_html=True)
            st.subheader("📊 Semáforo de Control Integral")
            
            df_mostrar = df_consolidado_rango.copy()
            df_mostrar["% Cumplimiento Rango"] = (df_mostrar["Venta Acumulada Rango ($)"] / (df_mostrar["Presupuesto Mes ($)"] / total_dias_mes * delta_dias) * 100)
            df_mostrar["% Crecimiento Dinamico"] = df_mostrar.apply(
                lambda row: ((row["Venta Acumulada Rango ($)"] - row["Venta Proporcional Año Pasado ($)"]) / row["Venta Proporcional Año Pasado ($)"] * 100) if row["Venta Proporcional Año Pasado ($)"] > 0 else 0.0,
                axis=1
            )
            df_mostrar["% Desperdicio"] = df_mostrar.apply(
                lambda row: (row["Desperdicio Acumulado (Unid)"] / row["Unidades Vendidas"] * 100) if row["Unidades Vendidas"] > 0 else 0.0,
                axis=1
            )
            
            col_crecimiento_titulo = f"% Crecimiento en Rango ({anio_sel} vs {anio_pasado})"
            df_formato = df_mostrar.copy()
            df_formato["Presupuesto Mes ($)"] = df_formato["Presupuesto Mes ($)"].apply(lambda x: f"${x:,.0f}")
            df_formato[col_vp_nombre] = df_formato[col_vp_nombre].apply(lambda x: f"${x:,.0f}")
            df_formato["Venta Proporcional Año Pasado ($)"] = df_formato["Venta Proporcional Año Pasado ($)"].apply(lambda x: f"${x:,.0f}")
            df_formato["Venta Acumulada Rango ($)"] = df_formato["Venta Acumulada Rango ($)"].apply(lambda x: f"${x:,.0f}")
            df_formato["Ticket Promedio ($)"] = df_formato["Ticket Promedio ($)"].apply(lambda x: f"${x:,.0f}")
            df_formato["Unidades Vendidas"] = df_formato["Unidades Vendidas"].apply(lambda x: f"{x:,}")
            df_formato["Desperdicio Acumulado (Unid)"] = df_formato["Desperdicio Acumulado (Unid)"].apply(lambda x: f"{x:,}")
            df_formato["% Desperdicio Fila"] = df_formato["% Desperdicio"].apply(lambda x: f"{x:.2f}%")
            df_formato["% Cumplimiento Rango Fila"] = df_mostrar["% Cumplimiento Rango"].apply(lambda x: f"{x:.2f}%")
            df_formato[col_crecimiento_titulo] = df_mostrar["% Crecimiento Dinamico"].apply(lambda x: f"{x:.2f}%")
            
            df_para_mostrar = df_formato[[
                "Almacén", "Presupuesto Mes ($)", col_vp_nombre, "Venta Proporcional Año Pasado ($)", 
                "Venta Acumulada Rango ($)", "Ticket Promedio ($)", "Unidades Vendidas", 
                "Desperdicio Acumulado (Unid)", "% Desperdicio Fila", "% Cumplimiento Rango Fila", col_crecimiento_titulo
            ]]

            def color_semaforo_integral(row):
                estilos = [''] * len(row)
                try:
                    val_desp = float(row['% Desperdicio Fila'].replace('%', '').strip())
                    idx_desp = row.index.get_loc('% Desperdicio Fila')
                    estilos[idx_desp] = 'background-color: #d4edda; color: #155724; font-weight: bold;' if val_desp <= 10.0 else 'background-color: #f8d7da; color: #721c24; font-weight: bold;'
                except:
                    pass
                try:
                    val_cump = float(row['% Cumplimiento Rango Fila'].replace('%', '').strip())
                    idx_cump = row.index.get_loc('% Cumplimiento Rango Fila')
                    if val_cump >= 100:
                        estilos[idx_cump] = 'background-color: #d4edda; color: #155724; font-weight: bold;'
                    elif val_cump >= 85:
                        estilos[idx_cump] = 'background-color: #fff3cd; color: #856404; font-weight: bold;'
                    else:
                        estilos[idx_cump] = 'background-color: #f8d7da; color: #721c24; font-weight: bold;'
                except:
                    pass
                try:
                    val_crec = float(row[col_crecimiento_titulo].replace('%', '').strip())
                    idx_crec = row.index.get_loc(col_crecimiento_titulo)
                    if val_crec > 0:
                        estilos[idx_crec] = 'background-color: #d4edda; color: #155724; font-weight: bold;'
                    elif val_crec < 0:
                        estilos[idx_crec] = 'background-color: #f8d7da; color: #721c24; font-weight: bold;'
                    else:
                        estilos[idx_crec] = 'background-color: #e2e3e5; color: #383d41; font-weight: bold;'
                except:
                    pass
                return estilos

            df_estilizado = df_para_mostrar.style.apply(color_semaforo_integral, axis=1)
            column_config_dict = {
                "Almacén": st.column_config.TextColumn("Almacén", width="medium"),
                "Presupuesto Mes ($)": st.column_config.TextColumn("Presupuesto Mes ($)", width="small"),
                col_vp_nombre: st.column_config.TextColumn(col_vp_nombre, width="small"),
                "Venta Proporcional Año Pasado ($)": st.column_config.TextColumn("Venta Prop. Año Pasado", width="small"),
                "Venta Acumulada Rango ($)": st.column_config.TextColumn("Venta Acum. Rango", width="small"),
                "Ticket Promedio ($)": st.column_config.TextColumn("Ticket Promedio", width="small"),
                "Unidades Vendidas": st.column_config.TextColumn("Unid. Vendidas", width="small"),
                "Desperdicio Acumulado (Unid)": st.column_config.TextColumn("Desperdicio", width="small"),
                "% Desperdicio Fila": st.column_config.TextColumn("% Desp.", width="small"),
                "% Cumplimiento Rango Fila": st.column_config.TextColumn("% Cump.", width="small"),
                col_crecimiento_titulo: st.column_config.TextColumn("% Crec.", width="small"),
            }
            st.dataframe(df_estilizado, use_container_width=True, column_config=column_config_dict, hide_index=True)

            output_integral = BytesIO()
            with pd.ExcelWriter(output_integral, engine='openpyxl') as writer:
                df_para_mostrar.to_excel(writer, index=False, sheet_name='Analisis_Integral_Zona4')
            excel_integral_bytes = output_integral.getvalue()
            
            st.download_button(
                label=f"📥 Exportar Análisis Integral a Excel (Rango {f_inicio} al {f_fin})",
                data=excel_integral_bytes,
                file_name=f"Analisis_Integral_Zona4_{f_inicio}_al_{f_fin}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        with tab_presupuestos:
            st.subheader(f"⚙️ Configuración de Presupuestos & Venta {anio_pasado} - {mes_sel_nombre} {anio_sel}")
            st.write(f"Modifica el presupuesto proyectado y registra la venta real obtenida en el mismo mes del año **{anio_pasado}**:")
            df_presup_edit = st.data_editor(df_presupuesto_activo, num_rows="fixed", hide_index=True, key=f"editor_presup_{anio_sel}_{mes_num}")
            
            col_p_save, col_p_dl, col_p_ul = st.columns(3)
            with col_p_save:
                if st.button("Guardar Configuración y Metas del Mes"):
                    guardar_presupuesto(anio_sel, mes_num, df_presup_edit)
                    st.success(f"¡Configuración de {mes_sel_nombre} {anio_sel} guardada de forma permanente!")
                    st.rerun()
            with col_p_dl:
                output_p = BytesIO()
                with pd.ExcelWriter(output_p, engine='openpyxl') as writer:
                    df_presupuesto_activo.to_excel(writer, index=False, sheet_name=f'Presupuestos_{mes_sel_nombre}')
                excel_p_bytes = output_p.getvalue()
                st.download_button(
                    label="📥 Descargar Plantilla Presupuestos",
                    data=excel_p_bytes,
                    file_name=f"Plantilla_Presupuestos_{mes_sel_nombre}_{anio_sel}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            with col_p_ul:
                arch_p = st.file_uploader("Subir Presupuestos diligenciados", type=["xlsx"], key="upload_presup")
                if arch_p:
                    try:
                        df_p_excel = pd.read_excel(arch_p)
                        if "Almacén" in df_p_excel.columns and "Presupuesto Mes ($)" in df_p_excel.columns:
                            guardar_presupuesto(anio_sel, mes_num, df_p_excel)
                            st.success("¡Presupuestos cargados y actualizados correctamente!")
                            st.rerun()
                        else:
                            st.error("El archivo Excel no tiene las columnas requeridas ('Almacén', 'Presupuesto Mes ($)').")
                    except Exception as e:
                        st.error(f"Error al leer el archivo: {e}")
            
        with tab_individual:
            dia_corte_ind = dia_corte_default
            def filtrar_por_rango(df, inicio, fin):
                df_copia = df.copy()
                df_copia["Fecha_Obj"] = pd.to_datetime(df_copia["Fecha y Día"].apply(lambda x: x.split(" ")[0])).dt.date
                return df_copia[(df_copia["Fecha_Obj"] >= inicio) & (df_copia["Fecha_Obj"] <= fin)]

            df_mes_corte_ind = filtrar_por_rango(df_mes_activo, primer_dia_mes, datetime(anio_sel, mes_num, min(dia_corte_ind, total_dias_mes)).date())

            st.subheader(f"Registro Diario por Tienda - {mes_sel_nombre} {anio_sel}")
            almacen_sel = st.selectbox("Selecciona el almacén a auditar / operar:", st.session_state.lista_almacenes_base)
            
            presupuesto_tienda = float(df_presupuesto_activo[df_presupuesto_activo["Almacén"] == almacen_sel]["Presupuesto Mes ($)"].values[0])
            meta_presup_tienda_fecha = (presupuesto_tienda / total_dias_mes) * dia_corte_ind
            venta_pasada_tienda = float(df_presupuesto_activo[df_presupuesto_activo["Almacén"] == almacen_sel][col_vp_nombre].values[0])
            venta_prop_pasada_tienda = (venta_pasada_tienda / total_dias_mes) * dia_corte_ind
            
            df_tienda_diario = df_mes_activo[df_mes_activo["Almacén"] == almacen_sel].copy()
            df_tienda_corte = df_mes_corte_ind[df_mes_corte_ind["Almacén"] == almacen_sel].copy()
            
            venta_tienda = df_tienda_corte["Venta Diaria ($)"].sum()
            trans_tienda = df_tienda_corte["Transacciones / Clientes"].sum()
            unid_tienda = df_tienda_corte["Unidades Vendidas"].sum()
            desp_tienda = df_tienda_corte["Desperdicio (Unid)"].sum()
            
            cump_tienda = (venta_tienda / presupuesto_tienda * 100) if presupuesto_tienda > 0 else 0.0
            cump_prop_tienda = (venta_tienda / meta_presup_tienda_fecha * 100) if meta_presup_tienda_fecha > 0 else 0.0
            crecimiento_tienda = ((venta_tienda - venta_prop_pasada_tienda) / venta_prop_pasada_tienda * 100) if venta_prop_pasada_tienda > 0 else 0.0
            ticket_tienda = (venta_tienda / trans_tienda) if trans_tienda > 0 else 0.0
            pct_desp_tienda = (desp_tienda / unid_tienda * 100) if unid_tienda > 0 else 0.0

            st.markdown(f"### 📍 Indicadores Exclusivos de: {almacen_sel} (Corte D-1 al Día {dia_corte_ind})")
            
            col_t1, col_t2, col_t3, col_t4, col_t5 = st.columns(5)
            with col_t1:
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">💰 VENTA ACUMULADA</div>
                        <div class="erp-value">${venta_tienda:,.0f}</div>
                        <div style="font-size: 11px; color: #0077b6; margin-top: 4px;">Meta a la Fecha: ${meta_presup_tienda_fecha:,.0f}</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_t2:
                color_cump_card_tienda = "#28a745" if cump_prop_tienda >= 100 else ("#ffc107" if cump_prop_tienda >= 85 else "#dc3545")
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">📈 CUMPLIMIENTO A LA FECHA</div>
                        <div class="erp-value" style="color: {color_cump_card_tienda};">{cump_prop_tienda:.2f}%</div>
                        <div style="font-size: 11px; color: #555; margin-top: 4px;">Cump. Mes Total: {cump_tienda:.2f}%</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_t3:
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">🎫 TICKET PROMEDIO</div>
                        <div class="erp-value">${ticket_tienda:,.0f}</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_t4:
                color_desp_card = "#28a745" if pct_desp_tienda <= 10.0 else "#dc3545"
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">⚠️ % DESPERDICIO</div>
                        <div class="erp-value" style="color: {color_desp_card};">{pct_desp_tienda:.2f}%</div>
                        <div style="font-size: 11px; color: #555; margin-top: 4px;">Acum: {desp_tienda:,} unid</div>
                    </div>
                """, unsafe_allow_html=True)
            with col_t5:
                color_cv_card = "#28a745" if crecimiento_tienda >= 0 else "#dc3545"
                st.markdown(f"""
                    <div class="erp-card">
                        <div class="erp-title">📊 CRECIMIENTO A LA FECHA</div>
                        <div class="erp-value" style="color: {color_cv_card};">{crecimiento_tienda:.2f}%</div>
                        <div style="font-size: 11px; color: #555; margin-top: 4px;">vs {anio_pasado} a la fecha</div>
                    </div>
                """, unsafe_allow_html=True)

            st.markdown("<br>", unsafe_allow_html=True)
            st.write(f"Ingresa **Venta Diaria**, **Transacciones / Clientes**, **Unidades Vendidas** y **Desperdicio (Unid)** día a día:")
            
            df_tienda_diario["Ticket Promedio ($)"] = df_tienda_diario.apply(
                lambda row: round(row["Venta Diaria ($)"] / row["Transacciones / Clientes"], 0) if row["Transacciones / Clientes"] > 0 else 0.0,
                axis=1
            )
            df_diario_edit = st.data_editor(df_tienda_diario, num_rows="fixed", hide_index=True, key=f"edit_diario_{almacen_sel}_{anio_sel}_{mes_num}")
            
            col_btn_save, col_btn_dl, col_btn_ul = st.columns(3)
            with col_btn_save:
                if st.button("Guardar Registro Diario de esta Tienda"):
                    df_diario_edit["Ticket Promedio ($)"] = df_diario_edit.apply(
                        lambda row: round(row["Venta Diaria ($)"] / row["Transacciones / Clientes"], 0) if row["Transacciones / Clientes"] > 0 else 0.0,
                        axis=1
                    )
                    df_mes_activo.loc[df_mes_activo["Almacén"] == almacen_sel] = df_diario_edit
                    guardar_mes(anio_sel, mes_num, df_mes_activo)
                    st.success(f"¡Registros de {almacen_sel} para {mes_sel_nombre} guardados de forma permanente!")
                    st.rerun()

            with col_btn_dl:
                output_tienda = BytesIO()
                with pd.ExcelWriter(output_tienda, engine='openpyxl') as writer:
                    df_mes_activo.to_excel(writer, index=False, sheet_name=f'Registros_{mes_sel_nombre}')
                excel_tienda_bytes = output_tienda.getvalue()
                st.download_button(
                    label=f"📥 Descargar Plantilla Registros",
                    data=excel_tienda_bytes,
                    file_name=f"Plantilla_Registros_{mes_sel_nombre}_{anio_sel}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            with col_btn_ul:
                arch_reg = st.file_uploader("Subir Registros Diarios en Excel", type=["xlsx"], key="upload_reg_diario")
                if arch_reg:
                    try:
                        df_reg_excel = pd.read_excel(arch_reg)
                        if "Fecha y Día" in df_reg_excel.columns and "Almacén" in df_reg_excel.columns:
                            df_reg_excel["Ticket Promedio ($)"] = df_reg_excel.apply(
                                lambda row: round(row["Venta Diaria ($)"] / row["Transacciones / Clientes"], 0) if row["Transacciones / Clientes"] > 0 else 0.0,
                                axis=1
                            )
                            guardar_mes(anio_sel, mes_num, df_reg_excel)
                            st.success("¡Registros diarios cargados y actualizados correctamente en el sistema!")
                            st.rerun()
                        else:
                            st.error("El archivo Excel no tiene las columnas obligatorias ('Fecha y Día', 'Almacén').")
                    except Exception as e:
                        st.error(f"Error al leer el archivo: {e}")

    else:
        dia_corte_admin = dia_corte_default
        def filtrar_por_rango(df, inicio, fin):
            df_copia = df.copy()
            df_copia["Fecha_Obj"] = pd.to_datetime(df_copia["Fecha y Día"].apply(lambda x: x.split(" ")[0])).dt.date
            return df_copia[(df_copia["Fecha_Obj"] >= inicio) & (df_copia["Fecha_Obj"] <= fin)]

        primer_dia_mes = datetime(anio_sel, mes_num, 1).date()
        df_mes_corte_admin = filtrar_por_rango(df_mes_activo, primer_dia_mes, datetime(anio_sel, mes_num, min(dia_corte_admin, total_dias_mes)).date())

        ced = st.session_state.cedula_actual
        tienda = st.session_state.db_admins[ced]["tienda"]
        
        presupuesto_tienda = float(df_presupuesto_activo[df_presupuesto_activo["Almacén"] == tienda]["Presupuesto Mes ($)"].values[0])
        meta_presup_tienda_fecha = (presupuesto_tienda / total_dias_mes) * dia_corte_admin
        venta_pasada_tienda = float(df_presupuesto_activo[df_presupuesto_activo["Almacén"] == tienda][col_vp_nombre].values[0])
        venta_prop_pasada_tienda = (venta_pasada_tienda / total_dias_mes) * dia_corte_admin
        
        df_mi_tienda_corte = df_mes_corte_admin[df_mes_corte_admin["Almacén"] == tienda].copy()
        df_mi_tienda_diario = df_mes_activo[df_mes_activo["Almacén"] == tienda].copy()
        
        v_tienda = df_mi_tienda_corte["Venta Diaria ($)"].sum()
        tr_tienda = df_mi_tienda_corte["Transacciones / Clientes"].sum()
        unid_tienda = df_mi_tienda_corte["Unidades Vendidas"].sum()
        d_tienda = df_mi_tienda_corte["Desperdicio (Unid)"].sum()
        
        c_tienda = (v_tienda / presupuesto_tienda * 100) if presupuesto_tienda > 0 else 0.0
        cump_prop_tienda = (v_tienda / meta_presup_tienda_fecha * 100) if meta_presup_tienda_fecha > 0 else 0.0
        crecimiento_tienda = ((v_tienda - venta_prop_pasada_tienda) / venta_prop_pasada_tienda * 100) if venta_prop_pasada_tienda > 0 else 0.0
        tk_tienda = (v_tienda / tr_tienda) if tr_tienda > 0 else 0.0
        pct_desp_tienda = (d_tienda / unid_tienda * 100) if unid_tienda > 0 else 0.0

        st.info(f"Tienda asignada: {tienda} | Periodo: {mes_sel_nombre} {anio_sel} (Corte D-1 al Día {dia_corte_admin})")
        
        col_ad1, col_ad2, col_ad3, col_ad4, col_ad5 = st.columns(5)
        with col_ad1:
            st.markdown(f"""
                <div class="erp-card">
                    <div class="erp-title">💰 VENTA ACUMULADA</div>
                    <div class="erp-value">${v_tienda:,.0f}</div>
                    <div style="font-size: 11px; color: #0077b6; margin-top: 4px;">Meta a la Fecha: ${meta_presup_tienda_fecha:,.0f}</div>
                </div>
            """, unsafe_allow_html=True)
        with col_ad2:
            color_cump_card_tienda = "#28a745" if cump_prop_tienda >= 100 else ("#ffc107" if cump_prop_tienda >= 85 else "#dc3545")
            st.markdown(f"""
                <div class="erp-card">
                    <div class="erp-title">📈 CUMPLIMIENTO A LA FECHA</div>
                    <div class="erp-value" style="color: {color_cump_card_tienda};">{cump_prop_tienda:.2f}%</div>
                    <div style="font-size: 11px; color: #555; margin-top: 4px;">Cump. Mes Total: {c_tienda:.2f}%</div>
                </div>
            """, unsafe_allow_html=True)
        with col_ad3:
            st.markdown(f"""
                <div class="erp-card">
                    <div class="erp-title">🎫 TICKET PROMEDIO</div>
                    <div class="erp-value">${tk_tienda:,.0f}</div>
                </div>
            """, unsafe_allow_html=True)
        with col_ad4:
            color_desp_card = "#28a745" if pct_desp_tienda <= 10.0 else "#dc3545"
            st.markdown(f"""
                <div class="erp-card">
                    <div class="erp-title">⚠️ % DESPERDICIO</div>
                    <div class="erp-value" style="color: {color_desp_card};">{pct_desp_tienda:.2f}%</div>
                    <div style="font-size: 11px; color: #555; margin-top: 4px;">Acum: {d_tienda:,} unid</div>
                </div>
            """, unsafe_allow_html=True)
        with col_ad5:
            color_cv_card = "#28a745" if crecimiento_tienda >= 0 else "#dc3545"
            st.markdown(f"""
                <div class="erp-card">
                    <div class="erp-title">📊 CRECIMIENTO A LA FECHA</div>
                    <div class="erp-value" style="color: {color_cv_card};">{crecimiento_tienda:.2f}%</div>
                    <div style="font-size: 11px; color: #555; margin-top: 4px;">vs {anio_pasado} a la fecha</div>
                </div>
            """, unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        st.write("Registra tus ventas, transacciones, unidades vendidas y desperdicio día a día:")
        
        df_mi_tienda_diario["Ticket Promedio ($)"] = df_mi_tienda_diario.apply(
            lambda row: round(row["Venta Diaria ($)"] / row["Transacciones / Clientes"], 0) if row["Transacciones / Clientes"] > 0 else 0.0,
            axis=1
        )
        df_diario_admin_edit = st.data_editor(df_mi_tienda_diario, num_rows="fixed", hide_index=True, key=f"admin_diario_{tienda}_{anio_sel}_{mes_num}")
        
        col_btn_save, col_btn_dl, col_btn_ul = st.columns(3)
        with col_btn_save:
            if st.button("Guardar Reporte Diario de la Tienda"):
                df_diario_admin_edit["Ticket Promedio ($)"] = df_diario_admin_edit.apply(
                    lambda row: round(row["Venta Diaria ($)"] / row["Transacciones / Clientes"], 0) if row["Transacciones / Clientes"] > 0 else 0.0,
                    axis=1
                )
                df_mes_activo.loc[df_mes_activo["Almacén"] == tienda] = df_diario_edit
                guardar_mes(anio_sel, mes_num, df_mes_activo)
                st.success("¡Reporte diario guardado de forma permanente en el sistema!")
                st.rerun()

        with col_btn_dl:
            output_tienda_adm = BytesIO()
            with pd.ExcelWriter(output_tienda_adm, engine='openpyxl') as writer:
                df_mi_tienda_diario.to_excel(writer, index=False, sheet_name='Mi_Tienda_Registro')
            excel_tienda_adm_bytes = output_tienda_adm.getvalue()
            st.download_button(
                label=f"📥 Descargar Plantilla Tienda",
                data=excel_tienda_adm_bytes,
                file_name=f"Plantilla_Registro_{tienda.split(' - ')[0]}_{mes_sel_nombre}_{anio_sel}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        with col_btn_ul:
            arch_reg_adm = st.file_uploader("Subir Plantilla Diligenciada", type=["xlsx"], key="upload_reg_adm")
            if arch_reg_adm:
                try:
                    df_reg_adm_excel = pd.read_excel(arch_reg_adm)
                    if "Fecha y Día" in df_reg_adm_excel.columns:
                        df_reg_adm_excel["Ticket Promedio ($)"] = df_reg_adm_excel.apply(
                            lambda row: round(row["Venta Diaria ($)"] / row["Transacciones / Clientes"], 0) if row["Transacciones / Clientes"] > 0 else 0.0,
                            axis=1
                        )
                        df_mes_activo.loc[df_mes_activo["Almacén"] == tienda] = df_reg_adm_excel
                        guardar_mes(anio_sel, mes_num, df_mes_activo)
                        st.success("¡Registro de la tienda actualizado correctamente!")
                        st.rerun()
                    else:
                        st.error("El archivo Excel no tiene la columna requerida ('Fecha y Día').")
                except Exception as e:
                    st.error(f"Error al leer el archivo: {e}")

# ---> MÓDULO DE AUDITORÍA EOR (CONECTADO AL EXCEL OFICIAL) <---
elif st.session_state.modulo_actual == "EOR":
    st.title("📋 Módulo de Auditoría EOR - Zona 4")
    st.write("Realiza la auditoría oficial sincronizada directamente con el archivo Excel maestro (**EOR SEPTIEMBRE 2026.xlsx**):")

    if not os.path.exists(EXCEL_EOR_MAESTRO):
        st.error(f"⚠️ No se encontró el archivo maestro '{EXCEL_EOR_MAESTRO}' en el directorio. Asegúrate de colocarlo en la misma carpeta del proyecto.")
    else:
        wb_maestro = openpyxl.load_workbook(EXCEL_EOR_MAESTRO, data_only=True)
        sheet_rev = wb_maestro[' Revisión de operaciones']

        # Extraer estructura de ítems desde la hoja ' Revisión de operaciones'
        preguntas_eor = []
        cat_actual = ""
        
        for r in range(10, sheet_rev.max_row + 1):
            val_cat = sheet_rev.cell(r, 1).value
            val_id = sheet_rev.cell(r, 2).value
            val_desc = sheet_rev.cell(r, 3).value
            
            if val_cat and val_cat.strip():
                cat_actual = val_cat.strip()
            
            if val_id and val_desc and str(val_id).startswith(("RSI", "FS-", "RC", "GS", "RE", "TR", "Pregunta", "MD", "ES", "LS")):
                preguntas_eor.append({
                    "Categoria": cat_actual if cat_actual else "General",
                    "Id": str(val_id).strip(),
                    "Pregunta": str(val_desc).strip()
                })

        with st.form("form_eor_maestro"):
            col_e1, col_e2, col_e3 = st.columns(3)
            with col_e1:
                fecha_auditoria = st.date_input("Fecha de Auditoría:", value=datetime.now().date())
            with col_e2:
                tienda_auditada = st.selectbox("Almacén a Auditar:", st.session_state.lista_almacenes_base)
            with col_e3:
                responsable_tienda = st.text_input("Responsable / Administrador:", value="Administrador de Tienda")
                auditor_nombre = st.text_input("Evaluador / Consultora:", value="Carolina Zamora")

            st.markdown("---")
            st.markdown("### 🔍 Evaluación por Estándares (One Inspire / EOR)")

            respuestas_usuario = []
            
            cat_grupo = ""
            for idx, item in enumerate(preguntas_eor):
                if item["Categoria"] != cat_grupo:
                    cat_grupo = item["Categoria"]
                    st.markdown(f"#### 📌 {cat_grupo}")
                
                st.markdown(f"**[{item['Id']}]** {item['Pregunta']}")
                c_val, c_obs = st.columns([1, 2])
                with c_val:
                    calif = st.selectbox("Cumplimiento:", ["Sí", "No", "N/A"], key=f"eor_val_{idx}")
                with c_obs:
                    obs = st.text_input("Observación / Justificación (Por qué cumple o no):", key=f"eor_obs_{idx}", placeholder="Escribe aquí el motivo o hallazgo...")
                st.markdown("")
                
                respuestas_usuario.append({
                    "Categoria": item["Categoria"],
                    "Id": item["Id"],
                    "Pregunta": item["Pregunta"],
                    "Calificacion": calif,
                    "Observacion": obs
                })

            submitted_eor_real = st.form_submit_button("💾 Guardar Auditoría y Actualizar Excel Maestro", use_container_width=True)
            
            if submitted_eor_real:
                try:
                    wb_mod = openpyxl.load_workbook(EXCEL_EOR_MAESTRO)
                    
                    codigo_tienda_str = tienda_auditada.split(" - ")[0].strip()
                    nombre_hoja_eor = f"EOR {codigo_tienda_str}"
                    nombre_hoja_pa = f"P.A. {codigo_tienda_str}"
                    
                    if nombre_hoja_eor not in wb_mod.sheetnames:
                        ws_nueva_eor = wb_mod.create_sheet(title=nombre_hoja_eor)
                    else:
                        ws_nueva_eor = wb_mod[nombre_hoja_eor]
                        
                    if nombre_hoja_pa not in wb_mod.sheetnames:
                        ws_nueva_pa = wb_mod.create_sheet(title=nombre_hoja_pa)
                    else:
                        ws_nueva_pa = wb_mod[nombre_hoja_pa]

                    escribir_celda_segura(ws_nueva_eor, 'C2', tienda_auditada)
                    escribir_celda_segura(ws_nueva_eor, 'C3', f"{responsable_tienda} - {auditor_nombre}")
                    escribir_celda_segura(ws_nueva_eor, 'C4', datetime.combine(fecha_auditoria, datetime.min.time()))

                    brechas = [r for r in respuestas_usuario if r["Calificacion"] == "No"]
                    
                    escribir_celda_segura(ws_nueva_pa, 'A1', "Plan de Acción — Mejora Continua EOR Zona 4")
                    escribir_celda_segura(ws_nueva_pa, 'A3', "Accion correctiva / Plan de accion")
                    escribir_celda_segura(ws_nueva_pa, 'B3', "Responsable")
                    escribir_celda_segura(ws_nueva_pa, 'C3', "Fecha de ejecucion prevista")
                    escribir_celda_segura(ws_nueva_pa, 'D3', "Comentarios / Hallazgos")
                    escribir_celda_segura(ws_nueva_pa, 'E3', "Fecha de cierre real")
                    
                    row_pa = 4
                    for b in brechas:
                        ws_nueva_pa.cell(row=row_pa, column=1, value=f"Atender hallazgo en [{b['Id']}]: {b['Pregunta']}")
                        ws_nueva_pa.cell(row=row_pa, column=2, value=responsable_tienda)
                        ws_nueva_pa.cell(row=row_pa, column=3, value="INMEDIATO / 5 DÍAS")
                        ws_nueva_pa.cell(row=row_pa, column=4, value=b['Observacion'])
                        ws_nueva_pa.cell(row=row_pa, column=5, value="")
                        row_pa += 1

                    wb_mod.save(EXCEL_EOR_MAESTRO)
                    st.success(f"¡Auditoría de {tienda_auditada} guardada con éxito! El archivo Excel '{EXCEL_EOR_MAESTRO}' ha sido actualizado en la pestaña '{nombre_hoja_eor}' y su Plan de Acción en '{nombre_hoja_pa}'.")
                
                except Exception as e:
                    st.error(f"Error al actualizar el archivo Excel: {e}")

# ---> MÓDULOS EN CONSTRUCCIÓN <---
elif st.session_state.modulo_actual == "Inventarios":
    st.title("📦 Módulo de Inventarios")
    st.info("🚧 Módulo en construcción. Aquí controlaremos el stock y los pedidos.")

elif st.session_state.modulo_actual == "Planilla":
    st.title("📅 Módulo de Planilla Semanal")
    st.info("🚧 Módulo en construcción. Aquí gestionaremos turnos y HeadCount.")

elif st.session_state.modulo_actual == "Previsiones":
    st.title("🔮 Módulo de Previsiones")
    st.info("🚧 Módulo en construcción. Aquí proyectaremos las metas y presupuestos.")